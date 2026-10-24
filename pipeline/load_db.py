"""Load O*NET 31.0 and the three exposure datasets into DuckDB, as separate tables.

    python pipeline/load_db.py

Exposure sources are kept apart on purpose: their disagreement becomes the
uncertainty band on every task. See docs/methodology.md for every rule below.

Task-level exposure E_t in [0, 1]:
  e_gpt4   GPT-4 rater label from "GPTs are GPTs":  E0 -> 0, E2 -> 0.5, E1 -> 1   (the paper's beta)
  e_human  Human rater label from the same dataset, same mapping
  e_aei    Anthropic Economic Index task penetration, ONLY where usage was observed
           (penetration > 0). Zero means "not seen in usage data", not "not exposed",
           so it is left out of the band rather than dragging the low end to 0.
  low/high = 0.5 * min/max of the available values; mid = (low + high) / 2.
  The 0.5 turns "exposed" into "share of the task's time AI could absorb": the paper
  defines exposure as cutting the task's time by at least half.
Tasks missing from every task-level source get the occupation's mean (flag imputed=1).
Occupation-level context: Felten AIOE as a percentile across occupations (shown, not
used in A_t, because it does not separate tasks inside a job).
"""
from __future__ import annotations

import csv
import re

import duckdb
import openpyxl

from common import CONFIG, DB_PATH, ONET_DIR, RAW

LABEL = {"E0": 0.0, "E1": 1.0, "E2": 0.5}

KEYWORD_OVERRIDES = [  # (regex on task text, capability) applied before the GWA vote
    (r"\btranslat", "translation"),
    (r"\binterpret(?:s|ing)? (?:for|between)|\bsign language", "translation"),
    (r"\btranscrib", "speech"),
    (r"\b(?:telephone|phone calls?|call cent)", "speech"),
    (r"\b(?:program(?:s|ming)? |source code|software code|debug|scripts?\b)", "code"),
    (r"\b(?:photograph|images?\b|video footage|x-rays?)", "vision"),
]


def norm_text(s: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()


def tsv(name: str) -> str:
    return str(ONET_DIR / name).replace("\\", "/")


def main() -> None:
    if DB_PATH.exists():
        DB_PATH.unlink()
    con = duckdb.connect(str(DB_PATH))
    rd = "delim='\t', header=true, quote='', escape='', all_varchar=true"

    con.execute(f"""
        CREATE TABLE occupations AS
        SELECT o."O*NET-SOC Code" AS code, o.Title AS title, o.Description AS description,
               TRY_CAST(z."Job Zone" AS INT) AS job_zone
        FROM read_csv('{tsv("Occupation Data.txt")}', {rd}) o
        LEFT JOIN read_csv('{tsv("Job Zones.txt")}', {rd}) z ON z."O*NET-SOC Code" = o."O*NET-SOC Code";
    """)
    con.execute(f"""
        CREATE TABLE tasks AS
        SELECT "O*NET-SOC Code" AS code, CAST("Task ID" AS INT) AS task_id, Task AS task,
               COALESCE(NULLIF("Task Type", ''), 'Unrated') AS task_type, "Date" AS updated
        FROM read_csv('{tsv("Task Statements.txt")}', {rd});
    """)
    con.execute(f"""
        CREATE TABLE task_importance AS
        SELECT "O*NET-SOC Code" AS code, CAST("Task ID" AS INT) AS task_id,
               CAST("Data Value" AS DOUBLE) AS im
        FROM read_csv('{tsv("Task Ratings.txt")}', {rd}) WHERE "Scale ID" = 'IM';
    """)
    con.execute(f"""
        CREATE TABLE dwa AS
        SELECT h."DWA Element ID" AS dwa_id, h."IWA Element ID" AS iwa_id,
               h."GWA Element ID" AS gwa_id, h."DWA Element Name" AS name
        FROM read_csv('{tsv("GWAs to IWAs to DWAs.txt")}', {rd}) h;
        CREATE TABLE task_dwas AS
        SELECT DISTINCT "O*NET-SOC Code" AS code, CAST("Task ID" AS INT) AS task_id,
               "DWA Element ID" AS dwa_id
        FROM read_csv('{tsv("Tasks to DWAs.txt")}', {rd});
        CREATE TABLE emerging AS
        SELECT "O*NET-SOC Code" AS code, Task AS task, Category AS category, "Date" AS updated
        FROM read_csv('{tsv("Emerging Tasks.txt")}', {rd});
        CREATE TABLE gwa_capability AS
        SELECT * FROM read_csv('{str(CONFIG / "gwa_capability_map.csv").replace(chr(92), "/")}', header=true, all_varchar=true);
    """)

    # --- GPTs are GPTs: task-level human and GPT-4 labels ---------------------------
    rows = []
    with open(RAW / "gpts_full_labelset.tsv", encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            try:
                tid = int(float(r["Task ID"]))
            except (TypeError, ValueError):
                continue
            rows.append((r["O*NET-SOC Code"], tid, LABEL.get(r["human_exposure_agg"]),
                         LABEL.get(r["gpt4_exposure"]), norm_text(r["Task"])))
    con.execute("CREATE TABLE exp_gpt (code VARCHAR, task_id INT, e_human DOUBLE, e_gpt4 DOUBLE, task_norm VARCHAR)")
    con.executemany("INSERT INTO exp_gpt VALUES (?,?,?,?,?)", rows)

    # --- Anthropic Economic Index: observed task penetration (text keyed) ------------
    aei = {}
    with open(RAW / "aei_task_penetration.csv", encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f):
            aei[norm_text(r["task"])] = float(r["penetration"])
    con.execute("CREATE TABLE exp_aei (task_norm VARCHAR, penetration DOUBLE)")
    con.executemany("INSERT INTO exp_aei VALUES (?,?)", list(aei.items()))
    aei_job = []
    with open(RAW / "aei_job_exposure.csv", encoding="utf-8", newline="") as f:
        for r in csv.DictReader(f):
            aei_job.append((r["occ_code"], float(r["observed_exposure"])))
    con.execute("CREATE TABLE aei_job (soc6 VARCHAR, observed_exposure DOUBLE)")
    con.executemany("INSERT INTO aei_job VALUES (?,?)", aei_job)

    # --- Felten AIOE: occupation-level, 6-digit SOC 2010 ----------------------------
    wb = openpyxl.load_workbook(RAW / "AIOE_DataAppendix.xlsx", read_only=True)
    ws = wb["Appendix A"]
    aioe = [(str(a), float(c)) for a, b, c in ws.iter_rows(min_row=2, values_only=True) if a and c is not None]
    con.execute("CREATE TABLE exp_aioe_raw (soc6 VARCHAR, aioe DOUBLE)")
    con.executemany("INSERT INTO exp_aioe_raw VALUES (?,?)", aioe)
    con.execute("""CREATE TABLE exp_aioe AS
        SELECT soc6, aioe, PERCENT_RANK() OVER (ORDER BY aioe) AS aioe_pct FROM exp_aioe_raw""")

    # --- Normalised task text for AEI join -----------------------------------------
    # Same rule as norm_text() above, written in SQL.
    con.execute("""ALTER TABLE tasks ADD COLUMN task_norm VARCHAR;
        UPDATE tasks SET task_norm = trim(regexp_replace(lower(task), '[^a-z0-9]+', ' ', 'g'));""")

    # --- Task capability: keyword override, else DWA -> GWA vote --------------------
    con.execute("""
        CREATE TABLE task_cap_vote AS
        SELECT td.code, td.task_id, g.capability, COUNT(*) AS n
        FROM task_dwas td JOIN dwa d USING (dwa_id) JOIN gwa_capability g ON g.gwa_id = d.gwa_id
        GROUP BY ALL;
        CREATE TABLE task_cap_gwa AS
        SELECT code, task_id, arg_max(capability, n * 100 + (CASE capability WHEN 'human' THEN 1 ELSE 0 END)) AS capability
        FROM task_cap_vote GROUP BY code, task_id;
    """)
    task_rows = con.execute("SELECT t.code, t.task_id, t.task, g.capability FROM tasks t LEFT JOIN task_cap_gwa g USING (code, task_id)").fetchall()
    caps = []
    for code, tid, text, gcap in task_rows:
        cap, method = gcap, "gwa_vote" if gcap else None
        for pat, c in KEYWORD_OVERRIDES:
            if re.search(pat, text, re.I):
                cap, method = c, "keyword"
                break
        if cap is None:
            cap, method = "human", "default_no_dwa"
        caps.append((code, tid, cap, method))
    con.execute("CREATE TABLE task_capability (code VARCHAR, task_id INT, capability VARCHAR, method VARCHAR)")
    con.executemany("INSERT INTO task_capability VALUES (?,?,?,?)", caps)

    # --- Combined exposure band ------------------------------------------------------
    con.execute("""
        CREATE TABLE task_exposure AS
        WITH g AS (
            SELECT t.code, t.task_id,
                   COALESCE(a.e_gpt4, b.e_gpt4) AS e_gpt4,
                   COALESCE(a.e_human, b.e_human) AS e_human,
                   CASE WHEN x.penetration > 0 THEN x.penetration END AS e_aei
            FROM tasks t
            LEFT JOIN exp_gpt a ON a.code = t.code AND a.task_id = t.task_id
            LEFT JOIN (SELECT task_norm, AVG(e_gpt4) e_gpt4, AVG(e_human) e_human FROM exp_gpt GROUP BY task_norm) b
                   ON a.task_id IS NULL AND b.task_norm = t.task_norm
            LEFT JOIN exp_aei x ON x.task_norm = t.task_norm
        ), band AS (
            SELECT *, LEAST(COALESCE(e_gpt4, 9), COALESCE(e_human, 9), COALESCE(e_aei, 9)) AS lo,
                      GREATEST(COALESCE(e_gpt4, -1), COALESCE(e_human, -1), COALESCE(e_aei, -1)) AS hi
            FROM g
        ), occmean AS (
            SELECT code, AVG(lo) AS mlo, AVG(hi) AS mhi FROM band WHERE lo < 9 GROUP BY code
        )
        -- TIME_SHARE = 0.5: an E1 label means AI "cuts the task's time by at least 50%",
        -- so a fully exposed task contributes E_t = 0.5 of its time, E2 contributes 0.25.
        SELECT b.code, b.task_id, b.e_gpt4, b.e_human, b.e_aei,
               0.5 * CASE WHEN b.lo < 9 THEN b.lo ELSE COALESCE(o.mlo, 0.5) END AS low,
               0.5 * CASE WHEN b.hi > -1 THEN b.hi ELSE COALESCE(o.mhi, 1.0) END AS high,
               CASE WHEN b.lo < 9 THEN 0 ELSE 1 END AS imputed
        FROM band b LEFT JOIN occmean o USING (code);
    """)

    for t, in con.execute("SELECT table_name FROM information_schema.tables ORDER BY 1").fetchall():
        n = con.execute(f'SELECT COUNT(*) FROM "{t}"').fetchone()[0]
        print(f"{t:22s} {n:>8,}")
    total, gpt, aei_n, imp, band = con.execute("""
        SELECT COUNT(*), COUNT(e_gpt4), COUNT(e_aei), SUM(imputed), ROUND(AVG(high - low), 3)
        FROM task_exposure""").fetchone()
    print(f"tasks={total:,} with_gpt_labels={gpt:,} with_aei_usage={aei_n:,} imputed={imp:,} mean_band={band}")
    print(con.execute("SELECT capability, COUNT(*) FROM task_capability GROUP BY 1 ORDER BY 2 DESC").fetchall())
    con.close()


if __name__ == "__main__":
    main()
